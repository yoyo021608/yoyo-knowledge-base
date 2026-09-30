from io import BytesIO

import pytest
from docx import Document as WordDocument

from server.infra.content import DocumentContentError
from server.infra.document_content import DefaultDocumentContentLoader


def test_extract_text_csv_and_docx() -> None:
    """常见资料格式应转换成保留结构的可检索文本。"""
    loader = DefaultDocumentContentLoader()
    text = loader.extract_file("notes.md", "# 标题\n正文".encode(), "text/markdown")
    csv = loader.extract_file("data.csv", "名称,值\nRAG,1".encode(), "text/csv")

    document = WordDocument()
    document.add_heading("知识库", level=1)
    document.add_paragraph("版本可追溯")
    payload = BytesIO()
    document.save(payload)
    word = loader.extract_file(
        "notes.docx",
        payload.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert "正文" in text.content
    assert "RAG | 1" in csv.content
    assert "知识库" in word.content and "版本可追溯" in word.content


def test_web_loader_rejects_private_network() -> None:
    """网页导入不能借助服务端访问本机或内网。"""
    with pytest.raises(DocumentContentError, match="内网"):
        DefaultDocumentContentLoader().fetch_web("http://127.0.0.1/private")
