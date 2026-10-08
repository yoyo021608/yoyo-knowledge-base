"""网页抓取与常见办公文件文本提取。"""

import csv
import io
import ipaddress
import socket
from html.parser import HTMLParser
from http.client import HTTPMessage
from pathlib import Path
from typing import IO
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener
from zipfile import BadZipFile

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from server.infra.content import DocumentContentError, ExtractedContent

MAX_REMOTE_BYTES = 5 * 1024 * 1024
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".json", ".yaml", ".yml"}


def _public_url(url: str) -> None:
    """拒绝本机、内网和保留地址，避免网页导入成为内网访问入口。"""
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise DocumentContentError("网页地址必须是有效的 HTTP(S) 地址")
    try:
        default_port = 443 if parsed.scheme == "https" else 80
        addresses = socket.getaddrinfo(parsed.hostname, parsed.port or default_port)
    except OSError as exc:
        raise DocumentContentError("网页地址无法解析") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if not ip.is_global:
            raise DocumentContentError("网页导入不能访问本机、内网或保留地址")


class _SafeRedirectHandler(HTTPRedirectHandler):
    """每次跳转都重新校验目标，防止公开地址跳转到内网。"""

    def redirect_request(
        self,
        req: Request,
        fp: IO[bytes],
        code: int,
        msg: str,
        headers: HTTPMessage,
        newurl: str,
    ) -> Request | None:
        target = urljoin(req.full_url, newurl)
        _public_url(target)
        return super().redirect_request(req, fp, code, msg, headers, target)


class _ReadableHtml(HTMLParser):
    """忽略脚本样式，只保留标题和页面可见文本。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title_parts: list[str] = []
        self.text_parts: list[str] = []
        self._ignored = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag in {"script", "style", "noscript", "svg"}:
            self._ignored += 1
        if tag == "title":
            self._in_title = True
        if tag in {"p", "div", "article", "section", "li", "br", "h1", "h2", "h3"}:
            self.text_parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self._ignored:
            self._ignored -= 1
        if tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        if self._ignored:
            return
        value = " ".join(data.split())
        if not value:
            return
        if self._in_title:
            self.title_parts.append(value)
        self.text_parts.append(value)

    def result(self, fallback_title: str) -> ExtractedContent:
        lines = [line.strip() for line in " ".join(self.text_parts).splitlines()]
        content = "\n".join(line for line in lines if line)
        if not content:
            raise DocumentContentError("网页中没有可录入的正文")
        return ExtractedContent(
            title=" ".join(self.title_parts).strip() or fallback_title,
            content=content,
        )


class DefaultDocumentContentLoader:
    """实现公开网页、文本、PDF、Word、CSV 和 Excel 的内容提取。"""

    def fetch_web(self, url: str) -> ExtractedContent:
        _public_url(url)
        request = Request(
            url,
            headers={"User-Agent": "yoyo-knowledge-base/0.1 (+document import)"},
        )
        try:
            with build_opener(_SafeRedirectHandler()).open(
                request, timeout=15
            ) as response:
                content_type = response.headers.get_content_type()
                if content_type not in {"text/html", "text/plain", "text/markdown"}:
                    raise DocumentContentError("网页响应不是可读取的文本内容")
                raw = response.read(MAX_REMOTE_BYTES + 1)
                if len(raw) > MAX_REMOTE_BYTES:
                    raise DocumentContentError("网页正文超过 5 MiB 限制")
                charset = response.headers.get_content_charset() or "utf-8"
                text = raw.decode(charset, errors="replace")
        except DocumentContentError:
            raise
        except (HTTPError, URLError, TimeoutError, OSError) as exc:
            raise DocumentContentError("网页抓取失败") from exc
        fallback = urlsplit(url).hostname or "网页资料"
        if content_type in {"text/plain", "text/markdown"}:
            return ExtractedContent(title=fallback, content=text.strip())
        parser = _ReadableHtml()
        parser.feed(text)
        return parser.result(fallback)

    def extract_file(
        self, name: str, content: bytes, content_type: str
    ) -> ExtractedContent:
        if not content:
            raise DocumentContentError("上传文件不能为空")
        if len(content) > MAX_UPLOAD_BYTES:
            raise DocumentContentError("上传文件超过 20 MiB 限制")
        suffix = Path(name).suffix.lower()
        title = Path(name).stem.strip() or name
        try:
            if suffix == ".csv":
                csv_rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
                return ExtractedContent(
                    title, "\n".join(" | ".join(row) for row in csv_rows)
                )
            if suffix in _TEXT_SUFFIXES or content_type.startswith("text/"):
                return ExtractedContent(title, content.decode("utf-8-sig"))
            if suffix == ".pdf" or content_type == "application/pdf":
                pages = [
                    page.extract_text() or ""
                    for page in PdfReader(io.BytesIO(content)).pages
                ]
                return ExtractedContent(title, "\n\n".join(pages))
            if suffix == ".docx":
                document = WordDocument(io.BytesIO(content))
                paragraphs = [paragraph.text for paragraph in document.paragraphs]
                tables = [
                    " | ".join(cell.text for cell in row.cells)
                    for table in document.tables
                    for row in table.rows
                ]
                return ExtractedContent(title, "\n".join([*paragraphs, *tables]))
            if suffix in {".xlsx", ".xlsm"}:
                workbook = load_workbook(
                    io.BytesIO(content), read_only=True, data_only=True
                )
                rows: list[str] = []
                for sheet in workbook.worksheets:
                    rows.append(f"# {sheet.title}")
                    rows.extend(
                        " | ".join("" if value is None else str(value) for value in row)
                        for row in sheet.iter_rows(values_only=True)
                    )
                return ExtractedContent(title, "\n".join(rows))
        except (
            BadZipFile,
            PackageNotFoundError,
            PdfReadError,
            UnicodeDecodeError,
            ValueError,
            OSError,
        ) as exc:
            raise DocumentContentError("文件内容损坏或编码不受支持") from exc
        raise DocumentContentError("支持 TXT、Markdown、PDF、DOCX、CSV 和 XLSX 文件")
